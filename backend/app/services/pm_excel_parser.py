from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import unicodedata
import zipfile
from xml.etree import ElementTree
from posixpath import normpath

from fastapi import HTTPException, status
from app.services.pm import calculate_budget_sale_amount


MAX_XLSX_BYTES = 15 * 1024 * 1024
MAX_SHEETS = 40
MAX_ROWS = 5000
MAX_COLUMNS = 100
MAX_CELL_TEXT = 4000
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

FIELD_ALIASES = {
    "codigo": ("codigo", "clave", "no", "numero", "item", "code"),
    "capitulo": ("capitulo", "capitulo padre", "grupo", "seccion"),
    "concepto": ("concepto", "conceptos", "actividad", "descripcion", "descripción", "partida", "trabajo"),
    "unidad": ("unidad", "un", "medida", "udm"),
    "cantidad": ("cantidad", "cant", "volumen", "qty"),
    "precio_unitario": ("precio unitario", "pu", "p.u.", "precio", "unitario"),
    "importe": ("importe", "total", "subtotal", "monto"),
    "cantidad_contratada": ("cantidad contratada", "contratado", "volumen contratado"),
    "avance_anterior": ("avance anterior", "anterior", "estimado anterior"),
    "esta_estimacion": ("esta estimacion", "esta estimación", "periodo", "importe actual"),
    "acumulado": ("acumulado", "avance acumulado"),
    "por_ejecutar": ("por ejecutar", "saldo", "pendiente"),
}


def _text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text[:MAX_CELL_TEXT]


def _decimal(value: object) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    text = _text(value).replace("$", "").replace(",", "").replace(" ", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _is_summary_row(code: str | None, concept: str, field_values: dict[str, object], raw: list[object]) -> bool:
    normalized = unicodedata.normalize("NFKD", concept.strip().casefold())
    normalized = "".join(character for character in normalized if not unicodedata.combining(character))
    normalized = re.sub(r"[^a-z0-9]+", " ", normalized).strip()
    summary_labels = {
        "contrato", "contrato original", "total", "subtotal", "gran total",
        "total general", "total contrato", "importe total", "resumen", "resumen general",
    }
    if normalized in summary_labels:
        return True
    # Totals without a label often only contain cached formula values. A row
    # without either an item code or concept cannot become an operable item.
    return not code and not concept and any(value not in (None, "") for value in raw)


def _json_value(value: object) -> str | float | int | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value if not isinstance(value, str) else value[:MAX_CELL_TEXT]
    if isinstance(value, Decimal):
        return str(value)
    return _text(value)


def suggest_mapping(headers: list[str]) -> dict[str, int]:
    normalized = []
    for header in headers:
        folded = unicodedata.normalize("NFKD", header.strip().lower())
        without_marks = "".join(character for character in folded if not unicodedata.combining(character))
        normalized.append(re.sub(r"\s+", " ", without_marks).replace(".", ""))
    result: dict[str, int] = {}
    for field, aliases in FIELD_ALIASES.items():
        for index, header in enumerate(normalized):
            if header and any(header == alias.lower().replace(".", "") for alias in aliases):
                result[field] = index
                break
    return result


def _find_header(rows: list[list[object]]) -> tuple[int | None, list[str], dict[str, int]]:
    best: tuple[int, int, list[str], dict[str, int]] | None = None
    for index, row in enumerate(rows[:30]):
        headers = [_text(value) for value in row]
        mapping = suggest_mapping(headers)
        score = len(set(mapping) & {"concepto", "unidad", "cantidad", "precio_unitario", "importe"})
        if score >= 2 and (best is None or score > best[0]):
            best = (score, index, headers, mapping)
    if not best:
        return None, [], {}
    return best[1], best[2], best[3]


def _frts_estimation_sheet(sheet_names: list[str]) -> tuple[str | None, int | None]:
    matches: list[tuple[int, str]] = []
    for name in sheet_names:
        match = re.fullmatch(r"E(\d+)", name.strip(), re.IGNORECASE)
        if match:
            matches.append((int(match.group(1)), name))
    return (max(matches)[1], max(matches)[0]) if matches else (None, None)


def _header_merges(data: bytes) -> dict[str, list[tuple[int, int, int, int]]]:
    """Read merge semantics without expanding every worksheet into memory."""
    from openpyxl.utils.cell import range_boundaries

    main_ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    with zipfile.ZipFile(BytesIO(data)) as archive:
        relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {item.attrib["Id"]: item.attrib["Target"] for item in relationships}
        workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        result = {}
        for sheet in workbook.findall(f"{main_ns}sheets/{main_ns}sheet"):
            target = targets.get(sheet.attrib.get(f"{rel_ns}id"), "")
            path = target.lstrip("/") if target.startswith("/") else normpath(f"xl/{target}")
            if not path.startswith("xl/worksheets/") or path not in archive.namelist():
                continue
            merges = []
            with archive.open(path) as stream:
                for _event, element in ElementTree.iterparse(stream, events=("end",)):
                    if element.tag == f"{main_ns}mergeCell":
                        bounds = range_boundaries(element.attrib["ref"])
                        if bounds[1] <= 30 and bounds[2] <= MAX_COLUMNS:
                            merges.append(bounds)
                    element.clear()
            result[sheet.attrib["name"]] = merges
        return result


def _frts_multilevel_header(rows: list[list[object]], merges: list[tuple[int, int, int, int]]):
    leaf_index = next((index for index, row in enumerate(rows[:30])
                       if {"cantidad", "precio_unitario"}.issubset(suggest_mapping([_text(value) for value in row]))), None)
    if leaf_index is None:
        return None
    width = max((len(row) for row in rows[:leaf_index + 1]), default=0)
    expanded = [[_text(row[col]) if col < len(row) else "" for col in range(width)] for row in rows[:leaf_index + 1]]
    for min_col, min_row, max_col, max_row in merges:
        if min_row > leaf_index + 1 or min_row > len(expanded):
            continue
        value = expanded[min_row - 1][min_col - 1]
        for row in range(min_row - 1, min(max_row, len(expanded))):
            for col in range(min_col - 1, min(max_col, width)):
                expanded[row][col] = value
    start = next((index for index in range(max(0, leaf_index - 3), leaf_index + 1)
                  if {"codigo", "concepto", "unidad"}.issubset(suggest_mapping(expanded[index]))), None)
    if start is None or start == leaf_index:
        return None
    headers = [" ".join(dict.fromkeys(expanded[row][col] for row in range(start, leaf_index + 1) if expanded[row][col])) for col in range(width)]
    mapping = suggest_mapping(expanded[start])
    groups = {
        "proyecto": "importe", "contrato": "importe", "contratado": "importe",
        "avance anterior": "avance_anterior", "este avance": "esta_estimacion",
        "esta estimacion": "esta_estimacion", "avance acumulado": "acumulado",
        "acumulado": "acumulado", "por ejecutar": "por_ejecutar",
    }
    for col in range(width):
        group = expanded[start][col].strip().casefold()
        leaf = expanded[leaf_index][col].strip().casefold().replace(".", "")
        if group in groups and leaf.startswith("importe"):
            mapping[groups[group]] = col
        if group in {"proyecto", "contrato", "contratado"}:
            if leaf == "cantidad": mapping["cantidad_contratada"] = col
            if leaf in {"pu", "precio unitario"}: mapping["precio_unitario"] = col
    if not {"codigo", "concepto", "unidad", "cantidad_contratada", "precio_unitario", "importe"}.issubset(mapping):
        return None
    return leaf_index, headers, mapping, list(range(start + 1, leaf_index + 2))


def parse_xlsx(data: bytes, filename: str, content_type: str | None = None) -> dict:
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos .xlsx sin macros.")
    if content_type is not None and content_type.strip().lower() != XLSX_CONTENT_TYPE:
        raise HTTPException(status_code=400, detail="El tipo de archivo no coincide con un libro XLSX.")
    if not data or len(data) > MAX_XLSX_BYTES:
        raise HTTPException(status_code=400, detail="El archivo está vacío o excede el límite de 15 MB.")
    if not data.startswith(b"PK\x03\x04"):
        raise HTTPException(status_code=400, detail="El archivo no tiene una firma XLSX válida.")
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            names = archive.namelist()
            entries = archive.infolist()
            if any("vbaproject.bin" in name.lower() for name in names):
                raise HTTPException(status_code=400, detail="No se aceptan archivos con macros.")
            if "[Content_Types].xml" not in names or not any(name.startswith("xl/workbook") for name in names):
                raise HTTPException(status_code=400, detail="El archivo no tiene un formato XLSX válido.")
            if sum(entry.file_size for entry in entries) > 100 * 1024 * 1024 or any(
                entry.compress_size and entry.file_size / entry.compress_size > 1000 for entry in entries
            ):
                raise HTTPException(status_code=400, detail="El archivo excede los límites de seguridad para lectura.")
        from openpyxl import load_workbook

        formula_book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        value_book = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="No se pudo leer el archivo. Verifica que sea un XLSX válido.") from exc

    try:
        sheet_names = list(formula_book.sheetnames)
        if not sheet_names or len(sheet_names) > MAX_SHEETS:
            raise HTTPException(status_code=400, detail="El archivo debe tener entre 1 y 40 hojas.")
        sheet_payload: list[dict] = []
        frts_book = any(name.strip().upper() == "CARATULA" for name in sheet_names)
        merges_by_sheet = _header_merges(data) if frts_book else {}
        count = 0
        for name in sheet_names:
            f_sheet = formula_book[name]
            v_sheet = value_book[name]
            if f_sheet.max_column and f_sheet.max_column > MAX_COLUMNS:
                raise HTTPException(status_code=400, detail="Una hoja supera el límite de 100 columnas.")
            rows: list[list[dict]] = []
            formula_rows = f_sheet.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLUMNS)
            cached_rows = v_sheet.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLUMNS)
            for row_no, (formula_row, cached_row) in enumerate(zip(formula_rows, cached_rows), 1):
                values: list[dict] = []
                for col_no, cell in enumerate(formula_row):
                    raw_value = cell.value
                    cached_value = cached_row[col_no].value if col_no < len(cached_row) else None
                    formula = raw_value if isinstance(raw_value, str) and raw_value.startswith("=") else None
                    displayed = cached_value if formula else raw_value
                    values.append({"value": _json_value(displayed), "formula": _json_value(formula)})
                while values and values[-1]["value"] in (None, "") and not values[-1]["formula"]:
                    values.pop()
                rows.append(values)
            if len(rows) > MAX_ROWS:
                raise HTTPException(status_code=400, detail="Una hoja supera el límite de 5,000 filas.")
            count += len(rows)
            if count > MAX_ROWS:
                raise HTTPException(status_code=400, detail="El archivo supera el límite total de 5,000 filas.")
            simple_rows = [[cell["value"] for cell in row] for row in rows]
            header_index, headers, mapping = _find_header(simple_rows)
            header_rows = [header_index + 1] if header_index is not None else []
            if frts_book and re.fullmatch(r"E\d+(?:-RF)?", name.strip(), re.IGNORECASE):
                multilevel = _frts_multilevel_header(simple_rows, merges_by_sheet.get(name, []))
                if multilevel:
                    header_index, headers, mapping, header_rows = multilevel
            if header_index is None:
                column_count = max((len(row) for row in simple_rows), default=0)
                headers = [f"Columna {index + 1}" for index in range(column_count)]
            sheet_payload.append({
                "name": name,
                "rows": rows,
                "header_row": header_index,
                "header_rows": header_rows,
                "headers": headers,
                "suggested_mapping": mapping,
                "mapping_required": header_index is None,
            })
        frts_sheet, frts_number = _frts_estimation_sheet(sheet_names)
        has_frts_cover = any(name.strip().upper() == "CARATULA" for name in sheet_names)
        if frts_sheet and has_frts_cover:
            selected = frts_sheet
            format_type = "frts"
        else:
            selected = max(sheet_payload, key=lambda item: len(item["suggested_mapping"]), default={}).get("name")
            format_type = "generic"
        selected_data = next((item for item in sheet_payload if item["name"] == selected), None)
        mapping = selected_data["suggested_mapping"] if selected_data else {}
        interpreted = interpret_sheet({**selected_data, "format_type": format_type}, mapping) if selected_data else []
        return {
            "format_type": format_type,
            "selected_sheet": selected,
            "estimation_number": frts_number if format_type == "frts" else None,
            "estimation_name": f"Estimación {frts_number} (importada)" if format_type == "frts" else "Estimación importada",
            "sheets": [{key: value for key, value in sheet.items() if key != "rows"} for sheet in sheet_payload],
            "sheet_rows": {sheet["name"]: sheet["rows"] for sheet in sheet_payload},
            "mapping_by_sheet": {sheet["name"]: sheet["suggested_mapping"] for sheet in sheet_payload},
            "selected_mapping": mapping,
            "preview_rows": interpreted,
        }
    finally:
        formula_book.close()
        value_book.close()


def interpret_sheet(sheet_meta: dict | None, mapping: dict[str, int]) -> list[dict]:
    if not sheet_meta:
        return []
    header_row = sheet_meta.get("header_row")
    rows = sheet_meta.get("rows", [])
    source_row_numbers = sheet_meta.get("source_row_numbers")
    result: list[dict] = []
    for offset, row in enumerate(rows):
        source_row = source_row_numbers[offset] if source_row_numbers and offset < len(source_row_numbers) else offset + 1
        if header_row is not None and source_row - 1 <= header_row:
            continue
        raw = [cell.get("value") for cell in row]
        formula_warnings = []
        for cell in row:
            formula = cell.get("formula")
            value = str(cell.get("value") or "")
            if formula and value.startswith("#"):
                formula_warnings.append("formula_error")
            elif formula and cell.get("value") is None:
                formula_warnings.append("formula_without_cached_value")
        field_values: dict[str, object] = {}
        for field, column in mapping.items():
            if isinstance(column, int) and 0 <= column < len(raw):
                field_values[field] = raw[column]
        concept = _text(field_values.get("concepto"))
        code = _text(field_values.get("codigo")) or None
        # FRTS calls its contractual quantity column "Cantidad contratada".
        # It is the budget quantity when a separate generic quantity mapping
        # is not supplied.
        qty = _decimal(field_values.get("cantidad"))
        if qty is None:
            qty = _decimal(field_values.get("cantidad_contratada"))
        price = _decimal(field_values.get("precio_unitario"))
        amount_excel = _decimal(field_values.get("importe"))
        if not concept and not code and all(value in (None, "") for value in raw):
            continue
        if _is_summary_row(code, concept, field_values, raw):
            result.append({
                "source_row": source_row,
                "row_type": "ignored",
                "code": code,
                "chapter": _text(field_values.get("capitulo")) or None,
                "concept": concept,
                "unit": _text(field_values.get("unidad")) or None,
                "quantity": str(qty) if qty is not None else None,
                "unit_price": str(price) if price is not None else None,
                "amount_excel": str(amount_excel) if amount_excel is not None else None,
                "amount_calculated": None,
                "contracted_quantity": str(_decimal(field_values.get("cantidad_contratada"))) if _decimal(field_values.get("cantidad_contratada")) is not None else None,
                "previous_progress": str(_decimal(field_values.get("avance_anterior"))) if _decimal(field_values.get("avance_anterior")) is not None else None,
                "this_estimate": str(_decimal(field_values.get("esta_estimacion"))) if _decimal(field_values.get("esta_estimacion")) is not None else None,
                "accumulated": str(_decimal(field_values.get("acumulado"))) if _decimal(field_values.get("acumulado")) is not None else None,
                "remaining": str(_decimal(field_values.get("por_ejecutar"))) if _decimal(field_values.get("por_ejecutar")) is not None else None,
                "warnings": list(dict.fromkeys([*formula_warnings, "summary_row"])),
                "raw_values": raw,
            })
            continue
        row_type = "item"
        is_frts_chapter = (
            sheet_meta.get("format_type") == "frts"
            and bool(code)
            and bool(re.fullmatch(r"[IVXLCDM]+", code, flags=re.IGNORECASE))
            and bool(concept)
        )
        if is_frts_chapter or (concept and qty is None and price is None):
            row_type = "chapter"
        warnings = list(dict.fromkeys(formula_warnings))
        if row_type == "item":
            if not concept:
                warnings.append("missing_description")
            if qty is None or qty <= 0:
                warnings.append("invalid_quantity")
            if not _text(field_values.get("unidad")):
                warnings.append("missing_unit")
            if price is None or price < 0:
                warnings.append("invalid_unit_price")
        calculated = calculate_budget_sale_amount(qty, price) if qty is not None and price is not None else None
        if amount_excel is not None and calculated is not None and abs(amount_excel - calculated) > Decimal("0.02"):
            warnings.append("amount_mismatch")
        estimate_amount = _decimal(field_values.get("esta_estimacion"))
        if estimate_amount is not None and calculated is not None and estimate_amount > calculated + Decimal("0.02"):
            warnings.append("estimate_exceeds_budget")
        result.append({
            "source_row": source_row,
            "row_type": row_type,
            "code": code,
            "chapter": _text(field_values.get("capitulo")) or None,
            "concept": concept,
            "unit": _text(field_values.get("unidad")) or None,
            "quantity": str(qty) if qty is not None else None,
            "unit_price": str(price) if price is not None else None,
            "amount_excel": str(amount_excel) if amount_excel is not None else None,
            "amount_calculated": str(calculated) if calculated is not None else None,
            "contracted_quantity": str(_decimal(field_values.get("cantidad_contratada"))) if _decimal(field_values.get("cantidad_contratada")) is not None else None,
            "previous_progress": str(_decimal(field_values.get("avance_anterior"))) if _decimal(field_values.get("avance_anterior")) is not None else None,
            "this_estimate": str(_decimal(field_values.get("esta_estimacion"))) if _decimal(field_values.get("esta_estimacion")) is not None else None,
            "accumulated": str(_decimal(field_values.get("acumulado"))) if _decimal(field_values.get("acumulado")) is not None else None,
            "remaining": str(_decimal(field_values.get("por_ejecutar"))) if _decimal(field_values.get("por_ejecutar")) is not None else None,
            "warnings": list(dict.fromkeys(warnings)),
            "raw_values": raw,
        })
    return result
