from __future__ import annotations

from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
import hashlib
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import time
from types import SimpleNamespace
import unittest

from fastapi import HTTPException
from starlette.datastructures import Headers, UploadFile
from openpyxl import Workbook
from sqlalchemy import create_engine, event, func, select
from sqlalchemy.orm import Session, sessionmaker
from unittest.mock import patch

from app.models import AuditLog, Empresa, EmpresaPMConfig, EmpresaUsuario, Plan, Usuario
from app.models.pm import PMEstimacion, PMEstimacionDetalle, PMPresupuesto, PMPresupuestoPartida, PMProyecto, PMTarea, PMProyectoLineaBase
from app.models.pm_imports import PMEstimacionEvidencia, PMExcelImportRow, PMExcelImportSession
from app.services.pm import PMContext
from app.services.pm_excel_import import (
    cancel_import,
    add_import_row,
    confirm_import,
    create_import_session,
    create_staged_evidence,
    get_import_session,
    update_import_details,
    update_import_mapping,
    update_import_row,
    update_staged_evidence,
    _summarize,
)
from app.services.pm_excel_parser import MAX_COLUMNS, MAX_ROWS, MAX_SHEETS, MAX_XLSX_BYTES, XLSX_CONTENT_TYPE, parse_xlsx, suggest_mapping


async def _collect_stream(iterator) -> bytes:
    chunks = []
    async for chunk in iterator:
        chunks.append(chunk)
    return b"".join(chunks)


class PMExcelImportTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.backend_dir = Path(__file__).resolve().parents[1]
        cls.temp_root = Path(tempfile.mkdtemp(prefix="pm-excel-import-"))
        cls.template_db_path = cls.temp_root / "template.db"
        env = dict(__import__("os").environ)
        env["DATABASE_URL"] = f"sqlite:///{cls.template_db_path.as_posix()}"
        result = subprocess.run(
            [sys.executable, "-m", "alembic", "upgrade", "head"],
            cwd=cls.backend_dir,
            env=env,
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode:
            raise RuntimeError(f"Alembic template upgrade failed:\n{result.stdout}\n{result.stderr}")

    @classmethod
    def tearDownClass(cls) -> None:
        shutil.rmtree(cls.temp_root, ignore_errors=True)

    def setUp(self) -> None:
        self.db_path = self.temp_root / f"{self._testMethodName}.db"
        shutil.copyfile(self.template_db_path, self.db_path)
        self.engine = create_engine(f"sqlite:///{self.db_path.as_posix()}", future=True)

        @event.listens_for(self.engine, "connect")
        def enable_sqlite_foreign_keys(connection, _record) -> None:
            cursor = connection.cursor()
            cursor.execute("PRAGMA foreign_keys=ON")
            cursor.close()

        self.session_factory = sessionmaker(bind=self.engine, autoflush=False, autocommit=False, class_=Session)
        self.db = self.session_factory()
        plan = Plan(code="basic", name="Basic", modules=["pm"])
        self.db.add(plan)
        self.db.flush()
        self.context_a, self.project_a = self._tenant("alpha", plan)
        self.context_b, _ = self._tenant("bravo", plan)
        self.db.commit()

    def tearDown(self) -> None:
        self.db.close()
        self.engine.dispose()
        if self.db_path.exists():
            self.db_path.unlink()

    def _tenant(self, suffix: str, plan: Plan):
        company = Empresa(
            name=f"Empresa {suffix}", slug=f"empresa-{suffix}", plan_code=plan.code,
            access_status="active", trial_ends_at=datetime.now(timezone.utc) + timedelta(days=30),
        )
        user = Usuario(
            email=f"{suffix}@example.com", full_name=f"Usuario {suffix}",
            password_hash="test-hash", is_active=True, is_superadmin=False,
        )
        self.db.add_all([company, user])
        self.db.flush()
        membership = EmpresaUsuario(empresa_id=company.id, usuario_id=user.id, role="admin", is_active=True)
        config = EmpresaPMConfig(
            empresa_id=company.id, pm_enabled=True, pm_tareas_enabled=True, pm_materiales_enabled=True,
            pm_tiempo_enabled=True, pm_templates_enabled=False, pm_comercial_enabled=False, pm_portal_enabled=True,
        )
        project = PMProyecto(
            empresa_id=company.id, nombre=f"Trabajo {suffix}", estatus="activo", prioridad="media",
            created_by=user.id, updated_by=user.id,
        )
        self.db.add_all([membership, config, project])
        self.db.flush()
        return PMContext(user=user, empresa_id=company.id, membership_role="admin", config=config), project

    @staticmethod
    def _workbook_bytes(*, frts: bool = False) -> bytes:
        workbook = Workbook()
        cover = workbook.active
        cover.title = "CARATULA" if frts else "Presupuesto"
        if frts:
            for name in ("E1", "E3", "E3-RF", "E2"):
                workbook.create_sheet(name)
            sheet = workbook["E3"]
        else:
            sheet = cover
        sheet.append(["Código", "Capítulo", "Concepto", "Unidad", "Cantidad", "Precio unitario", "Importe"])
        sheet.append(["", "", "Acabados", "", "", "", ""])
        sheet.append(["01", "Acabados", "Pintura", "m2", 10, 100, 1000])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_frts_selects_highest_exact_estimation_sheet(self) -> None:
        parsed = parse_xlsx(self._workbook_bytes(frts=True), "presupuesto.xlsx")
        self.assertEqual(parsed["format_type"], "frts")
        self.assertEqual(parsed["selected_sheet"], "E3")
        self.assertNotEqual(parsed["selected_sheet"], "E3-RF")

    @staticmethod
    def _multilevel_frts_bytes(column_offset: int = 0) -> bytes:
        workbook = Workbook()
        workbook.active.title = "CARATULA"
        sheet = workbook.create_sheet("E8")
        workbook.create_sheet("E8-RF")
        def put(row, col, value):
            sheet.cell(row, col + column_offset, value)
        for col, label in ((1, "No."), (2, "CONCEPTOS"), (3, "Unidad")):
            put(8, col, label)
            sheet.merge_cells(start_row=8, end_row=10, start_column=col + column_offset, end_column=col + column_offset)
        for first, last, label in ((4, 7, "PROYECTO"), (8, 10, "AVANCE ANTERIOR"),
                                   (11, 13, "ESTE AVANCE"), (14, 16, "AVANCE ACUMULADO"), (17, 19, "POR EJECUTAR")):
            put(8, first, label)
            sheet.merge_cells(start_row=8, end_row=9, start_column=first + column_offset, end_column=last + column_offset)
        for col, label in ((4, "Cantidad"), (5, "P.U."), (6, "(%)"), (7, "Importe($)"),
                           (8, "Cantidad"), (9, "(%)"), (10, "Importe($)"), (11, "Cantidad"),
                           (12, "(%)"), (13, "Importe($)"), (14, "Cantidad"), (15, "(%)"),
                           (16, "Importe($)"), (17, "Cantidad"), (18, "(%)"), (19, "Importe($)")):
            put(10, col, label)
        put(11, 1, "A"); put(11, 2, "CONTRATO"); put(11, 4, 1); put(11, 5, 33265)
        put(12, 1, "I"); put(12, 2, "Cimentacion")
        for index in range(29):
            row = index + 13
            code = "CI1" if index in (0, 1) else "CI3.1" if index == 4 else f"P{index}"
            quantity, price = ("0.5", "2.01") if index < 4 else ("2765", "7.79") if index == 4 else ("1", "580249.04") if index == 28 else ("1", "1000")
            put(row, 1, code); put(row, 2, f"Partida {index}"); put(row, 3, "KG")
            put(row, 4, quantity); put(row, 5, price)
            put(row, 7, str(Decimal(quantity) * Decimal(price)))
            if index == 4:
                put(row, 16, "22006.75"); put(row, 19, "-467.40")
            if index == 28:
                put(row, 13, "36241.55"); put(row, 16, "495808.04"); put(row, 19, "107445.02")
        put(42, 2, "TOTAL")
        output = BytesIO(); workbook.save(output); workbook.close()
        return output.getvalue()

    def test_multilevel_auto_mapping_and_persistence(self) -> None:
        expected = {"codigo": 0, "concepto": 1, "unidad": 2, "cantidad_contratada": 3,
                    "precio_unitario": 4, "importe": 6, "avance_anterior": 9,
                    "esta_estimacion": 12, "acumulado": 15, "por_ejecutar": 18}
        for offset in (0, 2):
            with self.subTest(column_offset=offset):
                parsed = parse_xlsx(self._multilevel_frts_bytes(offset), "synthetic.xlsx")
                self.assertEqual(parsed["selected_sheet"], "E8")
                self.assertEqual(parsed["selected_mapping"], {key: value + offset for key, value in expected.items()})
                self.assertEqual(next(s for s in parsed["sheets"] if s["name"] == "E8")["header_rows"], [8, 9, 10])
        created = create_import_session(self.db, self.context_a, self.project_a.id, "synthetic.xlsx", self._multilevel_frts_bytes())
        self.db.commit(); self.db.expire_all()
        reopened = get_import_session(self.db, self.context_a, created["id"])
        self.assertEqual(reopened["mapping"], expected)
        self.assertEqual(reopened["summary"]["chapters_count"], 1)
        self.assertEqual(reopened["summary"]["items_count"], 29)
        self.assertEqual(reopened["summary"]["total_recalculated"], "624792.43")
        self.assertEqual(reopened["summary"]["total_detected"], "624792.41")
        self.assertEqual(reopened["summary"]["difference"], "0.02")
        self.assertEqual(reopened["summary"]["reconciliation_status"], "redondeo")
        rows = {row["source_row"]: row for row in reopened["rows"]}
        self.assertFalse(rows[11]["include"]); self.assertFalse(rows[42]["include"])
        self.assertEqual(rows[12]["row_type"], "chapter")
        self.assertIn("duplicate_code", rows[13]["warnings"])
        self.assertIn("accumulated_exceeds_contracted", rows[17]["warnings"])

    def test_rounding_preview_confirm_and_readback_are_identical(self) -> None:
        created = create_import_session(self.db, self.context_a, self.project_a.id, "rounding.xlsx", self._multilevel_frts_bytes())
        rows = {row["source_row"]: row for row in created["rows"]}
        update_import_row(self.db, self.context_a, created["id"], rows[14]["id"], {"code": "CI1A"})
        updated = update_import_row(self.db, self.context_a, created["id"], rows[17]["id"], {"accumulated": "21539.35", "remaining": "0"})
        self.assertEqual(updated["summary"]["errors_count"], 0)
        result = confirm_import(self.db, self.context_a, created["id"], warnings_acknowledged=True)
        self.db.commit(); self.db.expire_all()
        budget = self.db.get(PMPresupuesto, result["budget_id"])
        self.assertEqual(budget.total_venta, Decimal(updated["summary"]["total_recalculated"]))
        self.assertEqual(budget.total_venta, Decimal("624792.43"))
        details = self.db.scalars(select(PMEstimacionDetalle).where(PMEstimacionDetalle.estimacion_id == result["estimation_id"])).all()
        self.assertEqual(sum(d.importe_periodo for d in details), Decimal("36241.55"))
        self.assertEqual(sum(d.importe_periodo for d in details), Decimal(updated["summary"]["estimate_total_recalculated"]))
        self.assertEqual(self.db.scalar(select(func.count(PMTarea.id))), 0)
        self.assertEqual(self.db.scalar(select(func.count(PMProyectoLineaBase.id))), 0)

    def test_genuine_difference_is_not_classified_as_rounding(self) -> None:
        result = _summarize([{"row_type": "item", "values": {"quantity": "1", "unit_price": "100", "amount_excel": "99.98"}}])
        self.assertEqual(result["difference"], "0.02")
        self.assertEqual(result["reconciliation_status"], "diferencia")

    def test_external_frts_gate48_when_explicitly_supplied(self) -> None:
        source = __import__("os").environ.get("PM_FRTS_SMOKE_PATH")
        if not source:
            self.skipTest("External FRTS is opt-in; it is never copied into the repository")
        created = create_import_session(self.db, self.context_a, self.project_a.id, Path(source).name, Path(source).read_bytes())
        self.assertEqual(created["selected_sheet"], "E8")
        self.assertEqual(created["summary"]["chapters_count"], 1)
        self.assertEqual(created["summary"]["items_count"], 29)
        self.assertEqual(created["summary"]["total_detected"], "624792.41")
        self.assertEqual(created["summary"]["total_recalculated"], "624792.43")
        self.assertEqual(created["summary"]["estimate_total_recalculated"], "36241.55")
        self.assertEqual(created["summary"]["accumulated_total"], "517814.79")
        self.assertEqual(created["summary"]["remaining_total"], "106977.62")
        self.assertEqual(created["summary"]["reconciliation_status"], "redondeo")
        duplicates = [row for row in created["rows"] if row["values"].get("code") == "CI1"]
        self.assertEqual(len(duplicates), 2)
        self.assertTrue(all("duplicate_code" in row["warnings"] for row in duplicates))
        self.assertTrue(any("formula_error" in row["warnings"] for row in created["rows"]))
        overrun = next(row for row in created["rows"] if row["values"].get("code") == "CI3.1")
        self.assertIn("accumulated_exceeds_contracted", overrun["warnings"])
        update_import_row(self.db, self.context_a, created["id"], duplicates[1]["id"], {"code": "CI1A"})
        updated = update_import_row(self.db, self.context_a, created["id"], overrun["id"], {"accumulated": "21539.35", "remaining": "0"})
        self.assertEqual(updated["summary"]["errors_count"], 0)
        self.db.commit()
        with self.session_factory() as reopened_db:
            reopened = get_import_session(reopened_db, self.context_a, created["id"])
            self.assertEqual(reopened["mapping"], created["mapping"])
            self.assertEqual(reopened["summary"]["items_count"], 29)
            result = confirm_import(reopened_db, self.context_a, created["id"], warnings_acknowledged=True)
            reopened_db.commit()
        with self.session_factory() as readback:
            budget = readback.get(PMPresupuesto, result["budget_id"])
            self.assertEqual(budget.total_venta, Decimal("624792.43"))
            self.assertEqual(budget.total_venta, Decimal(updated["summary"]["total_recalculated"]))
            details = readback.scalars(select(PMEstimacionDetalle).where(PMEstimacionDetalle.estimacion_id == result["estimation_id"])).all()
            self.assertEqual(sum(d.importe_periodo for d in details), Decimal("36241.55"))
            self.assertEqual(readback.scalar(select(func.count(PMTarea.id))), 0)
            self.assertEqual(readback.scalar(select(func.count(PMProyectoLineaBase.id))), 0)

    def test_evidence_calendar_date_create_update_clear_and_reopen(self) -> None:
        created = create_import_session(self.db, self.context_a, self.project_a.id, "evidence.xlsx", self._workbook_bytes())
        upload = SimpleNamespace(archivo_url="private", blob_path="test-only", filename="qa.jpg", content_type="image/jpeg", size_bytes=10)
        evidence = create_staged_evidence(self.db, self.context_a, project_id=self.project_a.id, session_id=created["id"], upload=upload,
                                         source_row=None, description="QA", location=None, evidence_date="2026-10-01")
        self.assertEqual(evidence["fecha_evidencia"], "2026-10-01")
        original_created = evidence["created_at"]
        update_staged_evidence(self.db, self.context_a, project_id=self.project_a.id, session_id=created["id"], evidence_id=evidence["id"], payload={"fecha_evidencia": "2026-10-02"})
        self.db.commit(); self.db.expire_all()
        reopened = get_import_session(self.db, self.context_a, created["id"])["evidences"][0]
        self.assertEqual(reopened["fecha_evidencia"], "2026-10-02")
        self.assertEqual(reopened["created_at"].replace(tzinfo=None), original_created.replace(tzinfo=None))
        for invalid in ("2026-02-30", "20261001", "2026-10-01T00:00:00Z"):
            with self.assertRaises(HTTPException) as raised:
                update_staged_evidence(self.db, self.context_a, project_id=self.project_a.id, session_id=created["id"], evidence_id=evidence["id"], payload={"fecha_evidencia": invalid})
            self.assertEqual(raised.exception.status_code, 400)
        cleared = update_staged_evidence(self.db, self.context_a, project_id=self.project_a.id, session_id=created["id"], evidence_id=evidence["id"], payload={"fecha_evidencia": None})
        self.assertIsNone(cleared["fecha_evidencia"])

    def test_summary_rows_are_ignored_until_user_changes_their_type(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Código", "Concepto", "Unidad", "Cantidad", "Precio unitario", "", "Importe"])
        sheet.append(["A", "CONTRATO", "LTE", 1, 100, None, 100])
        sheet.append([None, "TOTAL", None, None, None, None, 100])
        sheet.append([None, None, None, None, None, None, 100])
        sheet.append(["01", "Excavación", "m3", 1, 100, None, 100])
        output = BytesIO()
        workbook.save(output)

        created = create_import_session(self.db, self.context_a, self.project_a.id, "resumen.xlsx", output.getvalue())
        summary_rows = [row for row in created["rows"] if "summary_row" in row["warnings"]]
        self.assertEqual([row["source_row"] for row in summary_rows], [2, 3, 4])
        self.assertTrue(all(row["row_type"] == "ignored" and not row["include"] for row in summary_rows))

        included = update_import_row(self.db, self.context_a, created["id"], summary_rows[0]["id"], {
            "row_type": "item", "include": True,
        })
        opted_in = next(row for row in included["rows"] if row["id"] == summary_rows[0]["id"])
        self.assertTrue(opted_in["include"])
        self.assertNotIn("summary_row", opted_in["warnings"])

    def test_reconciliation_treats_sub_cent_rounding_as_reconciled(self) -> None:
        result = _summarize([{
            "row_type": "item", "include": True,
            "values": {"quantity": "1", "unit_price": "100", "amount_excel": "100.009"},
        }])
        self.assertEqual(result["reconciliation_status"], "conciliado")
        self.assertEqual(result["difference"], "-0.01")

    def test_mapping_preserves_source_row_numbers_across_blank_rows(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Presupuesto"
        sheet.append(["Código", "Concepto", "Unidad", "Cantidad", "Precio unitario", "Importe"])
        sheet.append([None])
        sheet.append(["A-3", "Partida después de fila vacía", "m2", 2, 10, 20])
        output = BytesIO()
        workbook.save(output)
        created = create_import_session(
            self.db, self.context_a, self.project_a.id, "presupuesto.xlsx", output.getvalue(),
        )
        updated = update_import_mapping(self.db, self.context_a, created["id"], selected_sheet="Presupuesto", mapping={
            "codigo": 0, "concepto": 1, "unidad": 2, "cantidad": 3, "precio_unitario": 4, "importe": 5,
        })
        item = next(row for row in updated["rows"] if row["source_row"] == 3)
        self.assertEqual(item["values"]["code"], "A-3")
        self.assertEqual(item["values"]["concept"], "Partida después de fila vacía")
        self.assertEqual(item["values"]["quantity"], "2")
        self.assertEqual(item["values"]["unit_price"], "10")

    def test_frts_manual_remapping_reclassifies_and_includes_contract_rows(self) -> None:
        workbook = Workbook()
        cover = workbook.active
        cover.title = "CARATULA"
        sheet = workbook.create_sheet("E8")
        workbook.create_sheet("E8-RF")
        for _ in range(10):
            sheet.append([None] * 19)
        sheet.append(["A", "CONTRATO", "LTE", 1, 33265, None, None] + [None] * 12)
        sheet.append(["I", "MURO DE CONTENCIÓN", None, None, None] + [None] * 14)

        for index in range(1, 30):
            code = "CI1" if index in {1, 2} else "CI3.1" if index == 3 else f"C{index:02d}"
            quantity = 2000 if code == "CI3.1" else 1
            unit_price = Decimal("7.79") if code == "CI3.1" else Decimal("1000")
            if index == 29:
                unit_price = Decimal("582212.41")
            values = [None] * 19
            values[0] = code
            values[1] = f"Partida {code}"
            values[2] = "KG" if code == "CI3.1" else "PZA"
            values[3] = quantity
            values[4] = float(unit_price)
            values[6] = float(Decimal(quantity) * unit_price)
            if code == "CI3.1":
                values[15] = 16047.40
                values[18] = -467.40
            if index == 29:
                values[12] = 36241.55
                values[15] = 501767.39
                values[18] = 107445.02
            sheet.append(values)
        output = BytesIO()
        workbook.save(output)

        created = create_import_session(self.db, self.context_a, self.project_a.id, "representativo-frts.xlsx", output.getvalue())
        mapping = {
            "codigo": 0, "concepto": 1, "unidad": 2, "cantidad_contratada": 3,
            "precio_unitario": 4, "importe": 6, "avance_anterior": 9,
            "esta_estimacion": 12, "acumulado": 15, "por_ejecutar": 18,
        }
        updated = update_import_mapping(self.db, self.context_a, created["id"], selected_sheet="E8", mapping=mapping)
        self.assertEqual(updated["summary"]["chapters_count"], 1)
        self.assertEqual(updated["summary"]["items_count"], 29)
        self.assertEqual(updated["summary"]["total_recalculated"], "624792.41")
        self.assertEqual(updated["summary"]["estimate_total_detected"], "36241.55")
        self.assertEqual(updated["summary"]["accumulated_total"], "517814.79")
        self.assertEqual(updated["summary"]["remaining_total"], "106977.62")
        self.assertEqual(updated["summary"]["reconciliation_status"], "conciliado")
        reopened = get_import_session(self.db, self.context_a, created["id"])
        self.assertEqual(reopened["summary"]["chapters_count"], 1)
        self.assertEqual(reopened["summary"]["items_count"], 29)

        rows = {row["source_row"]: row for row in updated["rows"]}
        contract_row = rows[11]
        chapter_row = rows[12]
        self.assertEqual(contract_row["row_type"], "ignored")
        self.assertFalse(contract_row["include"])
        self.assertEqual(chapter_row["row_type"], "chapter")
        self.assertTrue(chapter_row["include"])
        ci1_rows = [row for row in updated["rows"] if row["values"].get("code") == "CI1"]
        self.assertEqual(len(ci1_rows), 2)
        self.assertTrue(all("duplicate_code" in row["warnings"] for row in ci1_rows))
        overrun = next(row for row in updated["rows"] if row["values"].get("code") == "CI3.1")
        self.assertIn("accumulated_exceeds_contracted", overrun["warnings"])
        self.assertEqual(overrun["values"]["quantity"], "2000")

        excluded = update_import_row(self.db, self.context_a, created["id"], ci1_rows[0]["id"], {
            "include": False, "_include_explicit": True,
        })
        remapped = update_import_mapping(self.db, self.context_a, created["id"], selected_sheet="E8", mapping=mapping)
        excluded_ci1 = next(row for row in remapped["rows"] if row["id"] == ci1_rows[0]["id"])
        remaining_ci1 = next(row for row in remapped["rows"] if row["id"] == ci1_rows[1]["id"])
        self.assertFalse(excluded_ci1["include"])
        self.assertNotIn("duplicate_code", remaining_ci1["warnings"])

        manually_ignored = next(row for row in remapped["rows"] if row["values"].get("code") == "C04")
        update_import_row(self.db, self.context_a, created["id"], manually_ignored["id"], {
            "row_type": "ignored", "_row_type_explicit": True,
        })
        remapped_after_type_exclusion = update_import_mapping(self.db, self.context_a, created["id"], selected_sheet="E8", mapping=mapping)
        persisted_manual_exclusion = next(row for row in remapped_after_type_exclusion["rows"] if row["id"] == manually_ignored["id"])
        self.assertFalse(persisted_manual_exclusion["include"])

        overrun = next(row for row in remapped_after_type_exclusion["rows"] if row["values"].get("code") == "CI3.1")
        corrected = update_import_row(self.db, self.context_a, created["id"], overrun["id"], {
            "accumulated": "15580", "remaining": "0",
        })
        corrected_overrun = next(row for row in corrected["rows"] if row["id"] == overrun["id"])
        self.assertNotIn("accumulated_exceeds_contracted", corrected_overrun["warnings"])
        self.assertNotIn("negative_remaining", corrected_overrun["warnings"])
        reopened_corrected = get_import_session(self.db, self.context_a, created["id"])
        persisted_overrun = next(row for row in reopened_corrected["rows"] if row["id"] == overrun["id"])
        self.assertEqual(persisted_overrun["values"]["accumulated"], "15580")
        self.assertNotIn("accumulated_exceeds_contracted", persisted_overrun["warnings"])
        self.assertEqual(reopened_corrected["summary"]["errors_count"], 0)

    def test_formula_without_cached_value_is_preserved_as_warning_not_executed(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Concepto", "Unidad", "Cantidad", "Precio unitario", "Importe"])
        sheet.append(["Instalación", "serv", "=1+1", 500, "=D2*C2"])
        output = BytesIO()
        workbook.save(output)
        parsed = parse_xlsx(output.getvalue(), "formulas.xlsx")
        row = parsed["preview_rows"][0]
        self.assertIn("formula_without_cached_value", row["warnings"])
        self.assertIsNone(row["quantity"])
        self.assertEqual(parsed["sheet_rows"]["Sheet"][1][2]["formula"], "=1+1")

    def test_parser_handles_synthetic_1000_rows(self) -> None:
        for count in (100, 1000):
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Código", "Capítulo", "Concepto", "Unidad", "Cantidad", "Precio unitario", "Importe"])
            for index in range(count):
                sheet.append([f"{index + 1:04d}", "Capítulo A", f"Partida {index + 1}", "pza", 1, 10, 10])
            output = BytesIO()
            workbook.save(output)
            started = time.perf_counter()
            parsed = parse_xlsx(output.getvalue(), f"{count}-partidas.xlsx")
            parse_elapsed = time.perf_counter() - started
            mapping_started = time.perf_counter()
            for parsed_sheet in parsed["sheets"]:
                suggest_mapping(parsed_sheet["headers"])
            mapping_elapsed = time.perf_counter() - mapping_started
            serialization_started = time.perf_counter()
            __import__("json").dumps(parsed["preview_rows"], ensure_ascii=False)
            serialization_elapsed = time.perf_counter() - serialization_started
            self.assertEqual(len(parsed["preview_rows"]), count)
            print(
                f"Synthetic XLSX ({count} filas): parse/detección {parse_elapsed:.3f}s; "
                f"mapping adicional {mapping_elapsed:.6f}s; "
                f"serialización preview {serialization_elapsed:.3f}s"
            )

    def test_file_size_boundary_and_boundary_plus_one(self) -> None:
        valid = self._workbook_bytes()
        import zipfile
        exact_archive = BytesIO()
        filler_size = MAX_XLSX_BYTES
        filler = b"0123456789abcdef" * ((filler_size // 16) + 1)
        for _ in range(3):
            exact_archive = BytesIO()
            with zipfile.ZipFile(BytesIO(valid), "r") as source, zipfile.ZipFile(exact_archive, "w") as target:
                for member in source.infolist():
                    target.writestr(member, source.read(member.filename))
                target.writestr("xl/codex-boundary.bin", filler[:filler_size], compress_type=zipfile.ZIP_STORED)
            difference = MAX_XLSX_BYTES - len(exact_archive.getvalue())
            if difference == 0:
                break
            filler_size += difference
        exact_limit = exact_archive.getvalue()
        self.assertEqual(len(exact_limit), MAX_XLSX_BYTES)
        parsed = parse_xlsx(exact_limit, "limite.xlsx", XLSX_CONTENT_TYPE)
        self.assertEqual(parsed["format_type"], "generic")
        with self.assertRaises(HTTPException) as too_large:
            parse_xlsx(exact_limit + b"\0", "limite.xlsx", XLSX_CONTENT_TYPE)
        self.assertEqual(too_large.exception.status_code, 400)

    def test_row_limit_boundary_and_boundary_plus_one(self) -> None:
        def workbook_bytes(row_count: int) -> bytes:
            workbook = Workbook()
            sheet = workbook.active
            for index in range(row_count):
                sheet.append([f"Fila {index + 1}"])
            output = BytesIO()
            workbook.save(output)
            return output.getvalue()

        parse_xlsx(workbook_bytes(MAX_ROWS), "filas.xlsx")
        with self.assertRaises(HTTPException):
            parse_xlsx(workbook_bytes(MAX_ROWS + 1), "filas.xlsx")

    def test_column_limit_boundary_and_boundary_plus_one(self) -> None:
        def workbook_bytes(column_count: int) -> bytes:
            workbook = Workbook()
            sheet = workbook.active
            sheet.cell(row=1, column=column_count, value="Última")
            output = BytesIO()
            workbook.save(output)
            return output.getvalue()

        parse_xlsx(workbook_bytes(MAX_COLUMNS), "columnas.xlsx")
        with self.assertRaises(HTTPException):
            parse_xlsx(workbook_bytes(MAX_COLUMNS + 1), "columnas.xlsx")

    def test_sheet_limit_boundary_and_boundary_plus_one(self) -> None:
        def workbook_bytes(sheet_count: int) -> bytes:
            workbook = Workbook()
            for index in range(sheet_count - 1):
                workbook.create_sheet(f"Hoja {index + 2}")
            output = BytesIO()
            workbook.save(output)
            return output.getvalue()

        parse_xlsx(workbook_bytes(MAX_SHEETS), "hojas.xlsx")
        with self.assertRaises(HTTPException):
            parse_xlsx(workbook_bytes(MAX_SHEETS + 1), "hojas.xlsx")

    def test_mime_zip_structure_and_extension_are_checked(self) -> None:
        data = self._workbook_bytes()
        with self.assertRaises(HTTPException):
            parse_xlsx(data, "libro.xlsx", "application/pdf")
        with self.assertRaises(HTTPException):
            parse_xlsx(b"MZ" + b"\0" * 100, "ejecutable.xlsx")
        with self.assertRaises(HTTPException):
            parse_xlsx(data, "libro.xlsm")
        arbitrary_zip = BytesIO()
        import zipfile
        with zipfile.ZipFile(arbitrary_zip, "w") as archive:
            archive.writestr("readme.txt", "no es un libro")
        with self.assertRaises(HTTPException):
            parse_xlsx(arbitrary_zip.getvalue(), "renombrado.xlsx")

    def test_staging_changes_persist_after_close_and_reopen(self) -> None:
        created = create_import_session(self.db, self.context_a, self.project_a.id, "presupuesto.xlsx", self._workbook_bytes())
        self.db.commit()
        row = next(item for item in created["rows"] if item["row_type"] == "item")
        update_import_row(self.db, self.context_a, created["id"], row["id"], {"concept": "Partida corregida", "include": False})
        self.db.commit()
        session_id, row_id = created["id"], row["id"]
        self.db.close()
        self.db = self.session_factory()
        reopened = get_import_session(self.db, self.context_a, session_id)
        persisted = next(item for item in reopened["rows"] if item["id"] == row_id)
        self.assertEqual(persisted["values"]["concept"], "Partida corregida")
        self.assertFalse(persisted["include"])

    def test_confirm_failure_rolls_back_partial_product_records(self) -> None:
        created = create_import_session(self.db, self.context_a, self.project_a.id, "presupuesto.xlsx", self._workbook_bytes())
        self.db.commit()

        def fail_estimation_insert(_conn, _cursor, statement, _parameters, _context, _many):
            if statement.lstrip().upper().startswith("INSERT INTO PM_ESTIMACIONES"):
                raise RuntimeError("forced local failure")

        event.listen(self.engine, "before_cursor_execute", fail_estimation_insert)
        try:
            with self.assertRaises(RuntimeError):
                confirm_import(self.db, self.context_a, created["id"], warnings_acknowledged=True)
        finally:
            event.remove(self.engine, "before_cursor_execute", fail_estimation_insert)
            self.db.rollback()
        self.assertEqual(self.db.scalar(select(func.count(PMPresupuesto.id)).where(PMPresupuesto.proyecto_id == self.project_a.id)), 0)
        self.assertEqual(self.db.scalar(select(func.count(PMPresupuestoPartida.id)).where(PMPresupuestoPartida.proyecto_id == self.project_a.id)), 0)
        self.assertEqual(self.db.scalar(select(func.count(PMEstimacion.id)).where(PMEstimacion.proyecto_id == self.project_a.id)), 0)
        self.assertEqual(self.db.scalar(select(func.count(PMEstimacionDetalle.id)).where(PMEstimacionDetalle.proyecto_id == self.project_a.id)), 0)
        self.assertEqual(self.db.scalar(select(func.count(PMProyectoLineaBase.id)).where(PMProyectoLineaBase.proyecto_id == self.project_a.id)), 0)
        reopened = get_import_session(self.db, self.context_a, created["id"])
        self.assertIn(reopened["status"], {"review", "ready"})

    def test_source_download_is_tenant_scoped_and_never_returns_blob_reference(self) -> None:
        from app.api.routes.pm import download_budget_import_source_endpoint

        created = create_import_session(self.db, self.context_a, self.project_a.id, "origen.xlsx", self._workbook_bytes())
        row = self.db.get(__import__("app.models.pm_imports", fromlist=["PMExcelImportSession"]).PMExcelImportSession, created["id"])
        row.original_file_reference = f"{self.context_a.empresa_id}/private/source.xlsx"
        row.original_file_size = 100
        row.original_file_content_type = XLSX_CONTENT_TYPE
        self.db.commit()
        public = get_import_session(self.db, self.context_a, created["id"])
        self.assertNotIn("original_file_reference", public)
        self.assertTrue(public["source_file"]["available"])
        with patch("app.api.routes.pm.read_pm_blob", return_value=b"private xlsx") as read_blob:
            with self.assertRaises(HTTPException) as denied:
                download_budget_import_source_endpoint(self.project_a.id, created["id"], self.context_b, self.db)
            self.assertEqual(denied.exception.status_code, 404)
            read_blob.assert_not_called()
            response = download_budget_import_source_endpoint(self.project_a.id, created["id"], self.context_a, self.db)
            self.assertEqual(response.body, b"private xlsx")
            self.assertIn("no-store", response.headers["cache-control"])

    def test_source_file_survives_new_request_session_in_persistent_private_storage(self) -> None:
        from app.api.routes.pm import create_budget_import_endpoint, download_budget_import_source_endpoint

        storage_root = self.temp_root / "b"
        content = self._workbook_bytes()
        uploaded_paths: list[str] = []

        def local_blob_path(key: str) -> Path:
            if not key or key.startswith(("/", "\\")) or ".." in key.replace("\\", "/").split("/"):
                raise AssertionError("Unsafe storage key")
            root = (storage_root / "private").resolve()
            path = (root / f"{hashlib.sha256(key.encode('utf-8')).hexdigest()}.blob").resolve()
            if not path.is_relative_to(root):
                raise AssertionError("Storage key escaped the private test container")
            return path

        def local_upload(*, empresa_id, project_id, session_id, data, filename, content_type):
            from app.services.storage import build_pm_import_source_blob_path

            key = build_pm_import_source_blob_path(
                empresa_id=empresa_id, project_id=project_id, session_id=session_id,
            )
            path = local_blob_path(key)
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
            uploaded_paths.append(key)
            return key

        def local_read(*, blob_path, private=False):
            if not private:
                raise AssertionError("Original budget files must use private storage")
            return local_blob_path(blob_path).read_bytes()

        upload = UploadFile(
            filename="../../nombre raro; presupuesto ü.xlsx",
            file=BytesIO(content),
            headers=Headers({"content-type": XLSX_CONTENT_TYPE}),
        )
        with patch("app.api.routes.pm.upload_private_pm_import_file", side_effect=local_upload):
            result = __import__("asyncio").run(create_budget_import_endpoint(
                self.project_a.id, upload, self.context_a, self.db,
            ))

        project_id = self.project_a.id
        company_id = self.context_a.empresa_id
        user_id = self.context_a.user.id
        config_id = self.context_a.config.id
        other_company_id = self.context_b.empresa_id
        other_user_id = self.context_b.user.id
        other_config_id = self.context_b.config.id
        self.assertEqual(len(uploaded_paths), 1)
        persisted = self.db.get(PMExcelImportSession, result["id"])
        self.assertEqual(persisted.original_file_reference, uploaded_paths[0])
        self.assertEqual(persisted.file_hash, hashlib.sha256(content).hexdigest())
        self.assertEqual(persisted.original_file_size, len(content))
        self.assertEqual(persisted.original_file_content_type, XLSX_CONTENT_TYPE)

        # A separate DB session models the next authenticated HTTP request/process.
        self.db.close()
        next_request_db = self.session_factory()
        try:
            user = next_request_db.get(Usuario, user_id)
            config = next_request_db.get(EmpresaPMConfig, config_id)
            next_request_context = PMContext(
                user=user, empresa_id=company_id,
                membership_role="admin", config=config,
            )
            with patch("app.api.routes.pm.read_pm_blob", side_effect=local_read) as read_blob:
                response = download_budget_import_source_endpoint(
                    project_id, result["id"], next_request_context, next_request_db,
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.body, content)
                self.assertEqual(hashlib.sha256(response.body).hexdigest(), persisted.file_hash)
                self.assertEqual(response.media_type, XLSX_CONTENT_TYPE)
                self.assertIn("filename*=UTF-8''nombre%20raro%3B%20presupuesto%20%C3%BC.xlsx", response.headers["content-disposition"])
                self.assertEqual(response.headers["cache-control"], "private, no-store")
                self.assertEqual(response.headers["x-content-type-options"], "nosniff")
                read_blob.assert_called_once_with(blob_path=uploaded_paths[0], private=True)

            other_user = next_request_db.get(Usuario, other_user_id)
            other_config = next_request_db.get(EmpresaPMConfig, other_config_id)
            other_context = PMContext(
                user=other_user, empresa_id=other_company_id,
                membership_role="admin", config=other_config,
            )
            with patch("app.api.routes.pm.read_pm_blob", side_effect=AssertionError("Cross-tenant request reached storage")):
                with self.assertRaises(HTTPException) as denied:
                    download_budget_import_source_endpoint(
                        project_id, result["id"], other_context, next_request_db,
                    )
                self.assertEqual(denied.exception.status_code, 404)

            confirm_import(next_request_db, next_request_context, result["id"], warnings_acknowledged=True)
            next_request_db.commit()
            self.assertEqual(next_request_db.get(PMExcelImportSession, result["id"]).status, "imported")
            self.assertTrue(local_blob_path(uploaded_paths[0]).is_file())
            with patch("app.api.routes.pm.read_pm_blob", side_effect=local_read):
                imported_response = download_budget_import_source_endpoint(
                    project_id, result["id"], next_request_context, next_request_db,
                )
            self.assertEqual(imported_response.body, content)
        finally:
            next_request_db.close()

    def test_blob_delete_failure_restores_previously_deleted_blob(self) -> None:
        from app.api.routes.pm import _delete_blobs_with_compensation
        from app.services.storage import StorageOperationError

        with patch("app.api.routes.pm.delete_pm_blob", side_effect=[(True, b"image-bytes"), StorageOperationError()]), \
             patch("app.api.routes.pm.restore_pm_blob") as restore_blob:
            with self.assertRaises(StorageOperationError):
                _delete_blobs_with_compensation([
                    ("tenant/a.jpg", False, "image/jpeg"),
                    ("tenant/b.jpg", False, "image/jpeg"),
                ])
        restore_blob.assert_called_once_with(
            blob_path="tenant/a.jpg", private=False, data=b"image-bytes", content_type="image/jpeg",
        )

    def test_cancel_import_deletes_source_and_staged_evidence_but_keeps_audit(self) -> None:
        from app.api.routes.pm import cancel_budget_import_endpoint

        created = create_import_session(self.db, self.context_a, self.project_a.id, "origen.xlsx", self._workbook_bytes())
        session = self.db.get(PMExcelImportSession, created["id"])
        original_ref = f"{self.context_a.empresa_id}/private/{session.id}/source.xlsx"
        session.original_file_reference = original_ref
        session.original_file_size = 128
        session.original_file_content_type = XLSX_CONTENT_TYPE
        evidence = create_staged_evidence(
            self.db, self.context_a, project_id=self.project_a.id, session_id=session.id,
            upload=SimpleNamespace(archivo_url="https://example.test/photo.jpg", blob_path=f"{self.context_a.empresa_id}/photo.jpg", filename="photo.jpg", content_type="image/jpeg", size_bytes=64),
            source_row=None, description=None, location=None,
        )
        self.db.commit()
        with patch("app.api.routes.pm.delete_pm_blob", return_value=(True, b"blob-backup")) as delete_blob:
            result = cancel_budget_import_endpoint(session.id, self.context_a, self.db)
        self.assertEqual(result["status"], "cancelled")
        self.assertFalse(result["source_file"]["available"])
        self.assertEqual({call.kwargs["blob_path"] for call in delete_blob.call_args_list}, {
            original_ref, f"{self.context_a.empresa_id}/photo.jpg",
        })
        self.assertEqual(len(delete_blob.call_args_list), 2)
        self.assertEqual(self.db.scalar(select(func.count(PMExcelImportRow.id)).where(PMExcelImportRow.session_id == session.id)), 0)
        persisted_evidence = self.db.get(PMEstimacionEvidencia, evidence["id"])
        self.assertFalse(persisted_evidence.activo)
        self.assertIsNone(self.db.get(PMExcelImportSession, session.id).original_file_reference)
        audit = self.db.scalar(select(AuditLog).where(AuditLog.action == "pm.excel_import.cancel", AuditLog.entity_id == session.id))
        self.assertIsNotNone(audit)
        self.assertIn("file_hash", audit.metadata_json)

    def test_staged_evidence_can_be_reassigned_and_edited_within_its_session(self) -> None:
        created = create_import_session(self.db, self.context_a, self.project_a.id, "origen.xlsx", self._workbook_bytes())
        item_rows = self.db.scalars(select(PMExcelImportRow).where(
            PMExcelImportRow.session_id == created["id"], PMExcelImportRow.row_type == "item",
        ).order_by(PMExcelImportRow.source_row)).all()
        self.assertGreaterEqual(len(item_rows), 1)
        first_row = item_rows[0].source_row
        second_row = item_rows[-1].source_row
        evidence = create_staged_evidence(
            self.db, self.context_a, project_id=self.project_a.id, session_id=created["id"],
            upload=SimpleNamespace(archivo_url="data:image/png;base64,AA==", blob_path="mock/image.png", filename="image.png", content_type="image/png", size_bytes=32),
            source_row=first_row, description="Antes", location="Frente A",
        )
        updated = update_staged_evidence(self.db, self.context_a, project_id=self.project_a.id,
            session_id=created["id"], evidence_id=evidence["id"], payload={
                "source_row": second_row, "descripcion": "Después", "ubicacion": "Frente B",
            })
        self.assertEqual(updated["source_row"], second_row)
        self.assertEqual(updated["descripcion"], "Después")
        self.assertEqual(updated["ubicacion"], "Frente B")
        audit = self.db.scalar(select(AuditLog).where(
            AuditLog.action == "pm.excel_import.evidence_update", AuditLog.entity_id == evidence["id"],
        ))
        self.assertIsNotNone(audit)
        with self.assertRaises(HTTPException) as denied:
            update_staged_evidence(self.db, self.context_b, project_id=self.project_a.id,
                session_id=created["id"], evidence_id=evidence["id"], payload={"source_row": second_row})
        self.assertEqual(denied.exception.status_code, 404)

    def test_cancel_import_restores_blobs_when_database_commit_fails(self) -> None:
        from app.api.routes.pm import cancel_budget_import_endpoint
        from fastapi import HTTPException

        created = create_import_session(self.db, self.context_a, self.project_a.id, "origen.xlsx", self._workbook_bytes())
        session = self.db.get(PMExcelImportSession, created["id"])
        original_ref = f"{self.context_a.empresa_id}/private/{session.id}/source.xlsx"
        session.original_file_reference = original_ref
        session.original_file_content_type = XLSX_CONTENT_TYPE
        self.db.commit()
        with patch("app.api.routes.pm.delete_pm_blob", return_value=(True, b"source-bytes")), \
             patch("app.api.routes.pm.restore_pm_blob") as restore_blob, \
             patch.object(self.db, "commit", side_effect=RuntimeError("local forced commit failure")):
            with self.assertRaises(HTTPException) as failed:
                cancel_budget_import_endpoint(session.id, self.context_a, self.db)
        self.assertEqual(failed.exception.status_code, 500)
        restore_blob.assert_called_once_with(
            blob_path=original_ref, private=True, data=b"source-bytes", content_type=XLSX_CONTENT_TYPE,
        )
        self.db.expire_all()
        persisted = self.db.get(PMExcelImportSession, session.id)
        self.assertEqual(persisted.status, "review")
        self.assertEqual(persisted.original_file_reference, original_ref)

    def test_evidence_delete_removes_blob_and_is_tenant_scoped(self) -> None:
        from app.api.routes.pm import deactivate_estimation_evidence_endpoint

        created = create_import_session(self.db, self.context_a, self.project_a.id, "origen.xlsx", self._workbook_bytes())
        blob_path = f"{self.context_a.empresa_id}/photo.jpg"
        evidence = create_staged_evidence(
            self.db, self.context_a, project_id=self.project_a.id, session_id=created["id"],
            upload=SimpleNamespace(archivo_url="https://example.test/photo.jpg", blob_path=blob_path, filename="photo.jpg", content_type="image/jpeg", size_bytes=64),
            source_row=None, description=None, location=None,
        )
        self.db.commit()
        with patch("app.api.routes.pm.delete_pm_blob", return_value=(True, b"photo-bytes")) as delete_blob:
            with self.assertRaises(HTTPException) as denied:
                deactivate_estimation_evidence_endpoint(evidence["id"], self.context_b, self.db)
            self.assertEqual(denied.exception.status_code, 404)
            delete_blob.assert_not_called()
            result = deactivate_estimation_evidence_endpoint(evidence["id"], self.context_a, self.db)
        self.assertTrue(result["ok"])
        delete_blob.assert_called_once_with(blob_path=blob_path, private=False)
        self.assertFalse(self.db.get(PMEstimacionEvidencia, evidence["id"]).activo)

    def test_evidence_download_is_authenticated_scoped_and_streams_safe_image(self) -> None:
        import asyncio
        from app.api.routes.pm import download_pm_evidence_endpoint
        from app.services.storage import StorageBlobNotFoundError

        blob_path = f"{self.context_a.empresa_id}/pm/projects/{self.project_a.id}/documents/2026/10/evidence.jpg"
        evidence = PMEstimacionEvidencia(
            empresa_id=self.context_a.empresa_id,
            proyecto_id=self.project_a.id,
            url_archivo="https://storage.invalid/should-not-be-used.jpg",
            blob_path=blob_path,
            nombre_archivo='../../foto"\r\n evidencia ñ.jpg',
            mime_type="image/jpeg",
            size_bytes=11,
            created_by=self.context_a.user.id,
            activo=True,
        )
        self.db.add(evidence)
        self.db.commit()
        expected = b"private-jpeg"

        with patch("app.api.routes.pm.read_pm_blob", return_value=expected) as read_blob:
            response = download_pm_evidence_endpoint(evidence.id, self.context_a, self.db)
            body = asyncio.run(_collect_stream(response.body_iterator))
        self.assertEqual(body, expected)
        self.assertEqual(response.media_type, "image/jpeg")
        self.assertEqual(response.headers["cache-control"], "private, no-store")
        self.assertEqual(response.headers["x-content-type-options"], "nosniff")
        self.assertTrue(response.headers["content-disposition"].startswith("inline;"))
        self.assertIn("filename*=UTF-8''foto%20evidencia%20%C3%B1.jpg", response.headers["content-disposition"])
        self.assertNotIn("\r", response.headers["content-disposition"])
        self.assertNotIn("\n", response.headers["content-disposition"])
        read_blob.assert_called_once_with(blob_path=blob_path, private=False, require_private_access=True)

        with patch("app.api.routes.pm.read_pm_blob", side_effect=AssertionError("Cross-tenant read reached storage")):
            with self.assertRaises(HTTPException) as denied:
                download_pm_evidence_endpoint(evidence.id, self.context_b, self.db)
        self.assertEqual(denied.exception.status_code, 404)

        with patch("app.api.routes.pm.read_pm_blob", side_effect=StorageBlobNotFoundError()):
            with self.assertRaises(HTTPException) as missing_blob:
                download_pm_evidence_endpoint(evidence.id, self.context_a, self.db)
        self.assertEqual(missing_blob.exception.status_code, 404)

    def test_evidence_download_rejects_missing_and_manipulated_blob_references(self) -> None:
        from app.api.routes.pm import download_pm_evidence_endpoint

        with self.assertRaises(HTTPException) as missing:
            download_pm_evidence_endpoint("not-a-real-evidence-id", self.context_a, self.db)
        self.assertEqual(missing.exception.status_code, 404)

        evidence = PMEstimacionEvidencia(
            empresa_id=self.context_a.empresa_id,
            proyecto_id=self.project_a.id,
            url_archivo="https://storage.invalid/image.jpg",
            blob_path=f"{self.context_a.empresa_id}/pm/projects/{self.context_b.empresa_id}/private.jpg",
            nombre_archivo="image.jpg",
            mime_type="image/jpeg",
            size_bytes=5,
            created_by=self.context_a.user.id,
            activo=True,
        )
        self.db.add(evidence)
        self.db.commit()
        with patch("app.api.routes.pm.read_pm_blob", side_effect=AssertionError("Invalid key reached storage")):
            with self.assertRaises(HTTPException) as invalid_key:
                download_pm_evidence_endpoint(evidence.id, self.context_a, self.db)
        self.assertEqual(invalid_key.exception.status_code, 404)


    def test_import_confirmation_is_tenant_scoped_and_does_not_create_tasks_or_baseline(self) -> None:
        created = create_import_session(
            self.db, self.context_a, self.project_a.id, "presupuesto.xlsx", self._workbook_bytes(),
        )
        self.db.commit()
        self.assertEqual(created["selected_sheet"], "Presupuesto")
        self.assertEqual(created["summary"]["items_count"], 1)
        item_preview = next(row for row in created["rows"] if row["row_type"] == "item")
        self.assertEqual(item_preview["values"]["code"], "01")
        self.assertEqual(item_preview["values"]["chapter"], "Acabados")
        self.assertEqual(item_preview["values"]["unit_price"], "100")
        self.assertEqual(created["summary"]["total_recalculated"], "1000.00")

        with self.assertRaises(HTTPException) as cross_tenant:
            get_import_session(self.db, self.context_b, created["id"])
        self.assertEqual(cross_tenant.exception.status_code, 404)

        result = confirm_import(self.db, self.context_a, created["id"], warnings_acknowledged=True)
        self.db.commit()
        self.assertEqual(result["budget_status"], "borrador")
        self.assertEqual(result["tasks_created"], 0)
        self.assertFalse(result["baseline_created"])
        budget_items = self.db.scalars(select(PMPresupuestoPartida).where(PMPresupuestoPartida.proyecto_id == self.project_a.id)).all()
        self.assertTrue(budget_items, result)
        self.assertEqual(Decimal(budget_items[-1].precio_unitario_manual), Decimal("100"), [(row.tipo, row.precio_unitario_manual, row.precio_unitario, row.subtotal_venta) for row in budget_items])
        persisted_budget = self.db.get(PMPresupuesto, result["budget_id"])
        self.assertEqual(Decimal(result["budget_total"]), Decimal("1000.00"), (result, [(row.tipo, row.precio_unitario_manual, row.precio_unitario, row.subtotal_venta) for row in budget_items], persisted_budget.subtotal_venta, persisted_budget.total_venta))
        self.assertEqual(self.db.scalar(select(func.count(PMPresupuesto.id)).where(PMPresupuesto.proyecto_id == self.project_a.id)), 1)
        self.assertEqual(self.db.scalar(select(func.count(PMEstimacion.id)).where(PMEstimacion.proyecto_id == self.project_a.id)), 1)
        self.assertEqual(self.db.scalar(select(func.count(PMEstimacionDetalle.id)).where(PMEstimacionDetalle.empresa_id == self.context_a.empresa_id)), 1)
        self.assertEqual(self.db.scalar(select(func.count(PMPresupuestoPartida.id)).where(PMPresupuestoPartida.proyecto_id == self.project_a.id, PMPresupuestoPartida.tipo == "partida")), 1)
        self.assertEqual(self.db.scalar(select(func.count(PMTarea.id)).where(PMTarea.proyecto_id == self.project_a.id)), 0)
        self.assertEqual(self.db.scalar(select(func.count(PMProyectoLineaBase.id)).where(PMProyectoLineaBase.proyecto_id == self.project_a.id)), 0)
        persisted_session = self.db.get(PMExcelImportSession, created["id"])
        self.assertEqual(persisted_session.status, "imported")
        self.assertIsNotNone(self.db.scalar(select(AuditLog.id).where(
            AuditLog.action == "pm.excel_import.confirm", AuditLog.entity_id == created["id"],
        )))

        repeated = confirm_import(self.db, self.context_a, created["id"], warnings_acknowledged=True)
        self.assertTrue(repeated["already_imported"])
        self.assertEqual(self.db.scalar(select(func.count(PMPresupuesto.id)).where(PMPresupuesto.proyecto_id == self.project_a.id)), 1)

    def test_cross_tenant_import_actions_are_rejected(self) -> None:
        created = create_import_session(
            self.db, self.context_a, self.project_a.id, "presupuesto.xlsx", self._workbook_bytes(),
        )
        self.db.commit()
        item_row = next(row for row in created["rows"] if row["row_type"] == "item")
        operations = [
            lambda: update_import_details(self.db, self.context_b, created["id"], {"estimation_name": "No autorizado"}),
            lambda: update_import_row(self.db, self.context_b, created["id"], item_row["id"], {"concept": "No autorizado"}),
            lambda: confirm_import(self.db, self.context_b, created["id"], warnings_acknowledged=True),
            lambda: cancel_import(self.db, self.context_b, created["id"]),
            lambda: create_staged_evidence(
                self.db, self.context_b, project_id=self.project_a.id, session_id=created["id"],
                upload=None, source_row=None, description=None, location=None,
            ),
        ]
        for operation in operations:
            with self.subTest(operation=operation):
                with self.assertRaises(HTTPException) as denied:
                    operation()
                self.assertEqual(denied.exception.status_code, 404)

    def test_duplicate_rows_are_blocked_before_confirmation(self) -> None:
        created = create_import_session(
            self.db, self.context_a, self.project_a.id, "presupuesto.xlsx", self._workbook_bytes(),
        )
        self.db.commit()
        result = add_import_row(self.db, self.context_a, created["id"], {
            "row_type": "item", "include": True, "code": "01", "chapter": "Acabados",
            "concept": "Pintura", "unit": "m2", "quantity": "10", "unit_price": "100",
        })
        self.assertGreaterEqual(result["summary"]["errors_count"], 2)
        duplicate = next(row for row in result["rows"] if row["source_row"] >= 100000)
        self.assertIn("duplicate_code", duplicate["warnings"])
        self.assertIn("duplicate_row", duplicate["warnings"])
        with self.assertRaises(HTTPException) as invalid:
            confirm_import(self.db, self.context_a, created["id"], warnings_acknowledged=True)
        self.assertEqual(invalid.exception.status_code, 400)
        self.db.rollback()
        self.assertEqual(self.db.scalar(select(func.count(PMPresupuesto.id)).where(PMPresupuesto.proyecto_id == self.project_a.id)), 0)

    def test_accumulated_amount_is_compared_to_contract_amount_not_quantity(self) -> None:
        created = create_import_session(
            self.db, self.context_a, self.project_a.id, "presupuesto.xlsx", self._workbook_bytes(),
        )
        row = next(item for item in created["rows"] if item["row_type"] == "item")
        updated = update_import_row(self.db, self.context_a, created["id"], row["id"], {
            "quantity": "10", "unit_price": "100", "contracted_quantity": "10",
            "accumulated": "1000", "remaining": "0",
        })
        persisted = next(item for item in updated["rows"] if item["id"] == row["id"])
        self.assertNotIn("accumulated_exceeds_contracted", persisted["warnings"])
        self.assertEqual(updated["summary"]["contracted_total"], "1000.00")

    def test_formula_error_warning_survives_chapter_reclassification(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Código", "Concepto", "Unidad", "Cantidad", "Precio unitario", "Importe"])
        sheet.append(["I", "Muro de contención", "", 1, "=SUM(#REF!)", "=E2*D2"])
        output = BytesIO()
        workbook.save(output)
        created = create_import_session(
            self.db, self.context_a, self.project_a.id, "presupuesto.xlsx", output.getvalue(),
        )
        row = next(item for item in created["rows"] if item["row_type"] == "item")
        source_warnings = set(row["warnings"])
        updated = update_import_row(self.db, self.context_a, created["id"], row["id"], {"row_type": "chapter"})
        persisted = next(item for item in updated["rows"] if item["id"] == row["id"])
        self.assertEqual(persisted["row_type"], "chapter")
        self.assertTrue(source_warnings.intersection(persisted["warnings"]))

    def test_import_is_blocked_when_project_has_detailed_budget(self) -> None:
        created = create_import_session(
            self.db, self.context_a, self.project_a.id, "presupuesto.xlsx", self._workbook_bytes(),
        )
        self.db.commit()
        budget = PMPresupuesto(
            empresa_id=self.context_a.empresa_id, proyecto_id=self.project_a.id, nombre="Presupuesto existente",
            version=1, estatus="borrador", moneda="MXN", activo=True,
        )
        self.db.add(budget)
        self.db.flush()
        self.db.add(PMPresupuestoPartida(
            empresa_id=self.context_a.empresa_id, proyecto_id=self.project_a.id, presupuesto_id=budget.id,
            nombre="Partida existente", tipo="partida", unidad="pza", cantidad=1,
        ))
        self.db.commit()
        with self.assertRaises(HTTPException) as blocked:
            confirm_import(self.db, self.context_a, created["id"], warnings_acknowledged=True)
        self.assertEqual(blocked.exception.status_code, 409)
        self.assertIn("presupuesto detallado", blocked.exception.detail)
        self.db.rollback()
        self.assertEqual(self.db.scalar(select(func.count(PMPresupuesto.id)).where(PMPresupuesto.proyecto_id == self.project_a.id)), 1)


if __name__ == "__main__":
    unittest.main()
